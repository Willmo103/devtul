"""
Microsoft Word (.docx) document processor for DevTul.
Converts markdown text to professionally styled Word documents (.docx),
featuring frontmatter extraction, syntax-styled code blocks with background shading,
callouts, bumped header sizes (+2pt), tables, and inline formatting.
Adapted from .agents/docs/to_docx_processor.py.
"""

import io
from pathlib import Path
import re
from typing import Any, Dict, Optional, Tuple, Union

from bs4 import BeautifulSoup, NavigableString
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
import markdown
import yaml


class DocxProcessor:
    """
    Renders Markdown content into high-fidelity Microsoft Word (.docx) documents.
    """

    @staticmethod
    def apply_default_styles(doc: docx.Document) -> None:
        """Applies a consistent professional style, font, and size to the document."""
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        font.color.rgb = RGBColor(51, 51, 51)

        _format = style.paragraph_format
        _format.line_spacing = 1.15
        _format.space_after = Pt(6)

    @staticmethod
    def set_codeblock_style(
        paragraph: Any, fill_hex: str = "F6F8FA", border_hex: str = "D0D7DE"
    ) -> None:
        """Adds light grey background shading and left accent border to a codeblock paragraph."""
        pPr = paragraph._p.get_or_add_pPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        pPr.append(shd)
        pbdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="{border_hex}"/></w:pBdr>'
        )
        pPr.append(pbdr)

    @staticmethod
    def preprocess_fenced_code_blocks(md_text: str) -> str:
        """
        Preprocesses ```code``` blocks into explicit <pre><code> HTML blocks
        so Python-Markdown does not misinterpret them inside lists.
        """
        pattern = re.compile(
            r"^[ \t]*```(?P<lang>\w*)[ \t]*\r?\n(?P<code>.*?)\r?\n[ \t]*```[ \t]*$",
            re.MULTILINE | re.DOTALL,
        )

        def repl(match: re.Match) -> str:
            lang = match.group("lang").strip()
            code = match.group("code")
            lang_class = f' class="language-{lang}"' if lang else ""
            return f"\n<pre><code{lang_class}>\n{code}\n</code></pre>\n"

        return pattern.sub(repl, md_text)

    @staticmethod
    def extract_frontmatter(md_text: str) -> Tuple[Dict[str, Any], str]:
        """Extracts YAML frontmatter from top of markdown document if present."""
        pattern = re.compile(r"^---\s*\r?\n(.*?)\r?\n---\s*\r?\n", re.DOTALL)
        match = pattern.match(md_text)
        if not match:
            return {}, md_text

        fm_raw = match.group(1)
        body_text = md_text[match.end():]

        try:
            fm_data = yaml.safe_load(fm_raw)
            if isinstance(fm_data, dict):
                return fm_data, body_text
        except Exception:
            pass
        return {}, md_text

    @classmethod
    def render_frontmatter(cls, doc: docx.Document, fm_data: Dict[str, Any]) -> None:
        """Renders YAML frontmatter metadata as a styled header section at top of docx."""
        if not fm_data:
            return

        fm = dict(fm_data)
        title = (
            fm.pop("Title", None)
            or fm.pop("title", None)
            or fm.pop("Name", None)
            or fm.pop("name", None)
        )
        if title:
            p = doc.add_heading(str(title), level=1)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            for run in p.runs:
                run.font.size = Pt(20)

        if fm:
            table = doc.add_table(rows=len(fm), cols=2)
            table.style = "Table Grid"
            table.autofit = True

            for idx, (k, v) in enumerate(fm.items()):
                row = table.rows[idx]
                cell_key = row.cells[0]
                cell_val = row.cells[1]

                cell_key.text = str(k).replace("_", " ").capitalize()
                cell_val.text = str(v)

                p_k = cell_key.paragraphs[0]
                p_k.paragraph_format.space_after = Pt(2)
                if p_k.runs:
                    p_k.runs[0].bold = True
                    p_k.runs[0].font.size = Pt(9.5)

                p_v = cell_val.paragraphs[0]
                p_v.paragraph_format.space_after = Pt(2)
                if p_v.runs:
                    p_v.runs[0].font.size = Pt(9.5)

                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>')
                cell_key._tc.get_or_add_tcPr().append(shd)

            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    @classmethod
    def render_code_block(
        cls, doc: docx.Document, element: Any, left_indent_inches: float = 0.2
    ) -> None:
        """Renders a code block (<pre> or <code> block) as styled Consolas paragraphs with background shading."""
        code_text = element.get_text()

        lines = code_text.splitlines()
        if lines and lines[0].strip().lower() in [
            "python",
            "cmd",
            "bash",
            "sh",
            "json",
            "yaml",
            "yml",
            "ps1",
            "powershell",
            "html",
            "css",
            "javascript",
            "js",
            "ts",
        ]:
            lines = lines[1:]

        while lines and not lines[0].strip():
            lines.pop(0)
        while lines and not lines[-1].strip():
            lines.pop()

        if not lines:
            return

        for idx, line in enumerate(lines):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2) if idx == 0 else Pt(0)
            p.paragraph_format.space_after = Pt(2) if idx == len(lines) - 1 else Pt(0)
            p.paragraph_format.left_indent = Inches(left_indent_inches)
            cls.set_codeblock_style(p)
            run = p.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(36, 41, 46)

    @classmethod
    def add_inline_formatting(cls, doc: docx.Document, paragraph: Any, element: Any) -> None:
        """Processes child nodes of an HTML element, preserving inline code, bold, italic, strikethrough, and links."""
        for child in element.contents:
            if isinstance(child, NavigableString):
                paragraph.add_run(str(child))
            elif child.name == "code":
                code_text = child.get_text()
                if "\n" in code_text or code_text.startswith(
                    ("python\n", "cmd\n", "bash\n", "# ", ":: ")
                ):
                    cls.render_code_block(doc, child, left_indent_inches=0.4)
                else:
                    run = paragraph.add_run(code_text)
                    run.font.name = "Consolas"
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(199, 37, 78)
            elif child.name in ["strong", "b"]:
                run = paragraph.add_run(child.get_text())
                run.bold = True
            elif child.name in ["em", "i"]:
                run = paragraph.add_run(child.get_text())
                run.italic = True
            elif child.name in ["del", "s", "strike"]:
                run = paragraph.add_run(child.get_text())
                run.font.strike = True
            elif child.name == "a":
                run = paragraph.add_run(child.get_text())
                run.font.color.rgb = RGBColor(3, 102, 214)
                run.underline = True
            elif child.name == "br":
                paragraph.add_run("\n")
            elif child.name == "pre":
                cls.render_code_block(doc, child, left_indent_inches=0.4)
            elif child.name == "p":
                cls.add_inline_formatting(doc, paragraph, child)
            else:
                paragraph.add_run(child.get_text())

    @classmethod
    def process_list_item(
        cls, doc: docx.Document, li: Any, list_style: str, indent_level: int = 1
    ) -> None:
        """Processes a single <li> element, separating text, nested code blocks, and sublists."""
        pres = li.find_all(["pre"], recursive=False)
        sublists = li.find_all(["ul", "ol"], recursive=False)

        p = doc.add_paragraph(style=list_style)
        p.paragraph_format.left_indent = Inches(0.25 * indent_level)

        for child in li.contents:
            if isinstance(child, NavigableString):
                p.add_run(str(child))
            elif child.name == "p":
                cls.add_inline_formatting(doc, p, child)
            elif child.name not in ["pre", "ul", "ol"]:
                soup_child = BeautifulSoup(str(child), "html.parser")
                target = soup_child.body if soup_child.body else soup_child
                cls.add_inline_formatting(doc, p, target)

        for pre in pres:
            cls.render_code_block(doc, pre, left_indent_inches=0.25 * indent_level + 0.25)

        for sublist in sublists:
            sub_style = "List Bullet" if sublist.name == "ul" else "List Number"
            for sub_li in sublist.find_all("li", recursive=False):
                cls.process_list_item(doc, sub_li, sub_style, indent_level + 1)

    @classmethod
    def html_to_docx(cls, html_content: str, doc: docx.Document) -> None:
        """Parses HTML elements from markdown and renders them into the docx document."""
        soup = BeautifulSoup(html_content, "html.parser")

        for element in soup.children:
            if isinstance(element, NavigableString):
                continue

            if element.name == "h1":
                p = doc.add_heading(element.get_text(), level=1)
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(6)
                for run in p.runs:
                    run.font.size = Pt(20)
            elif element.name == "h2":
                p = doc.add_heading(element.get_text(), level=2)
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(16)
            elif element.name == "h3":
                p = doc.add_heading(element.get_text(), level=3)
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(14)
            elif element.name == "h4":
                p = doc.add_heading(element.get_text(), level=4)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(12)
            elif element.name == "p":
                p = doc.add_paragraph()
                cls.add_inline_formatting(doc, p, element)
            elif element.name in ["ul", "ol"]:
                list_style = "List Bullet" if element.name == "ul" else "List Number"
                for li in element.find_all("li", recursive=False):
                    cls.process_list_item(doc, li, list_style, indent_level=1)
            elif element.name == "blockquote":
                p = doc.add_paragraph(style="Intense Quote")
                cls.add_inline_formatting(doc, p, element)
            elif element.name in ["pre", "div"] and (
                element.name == "pre" or "code" in element.get("class", [])
            ):
                cls.render_code_block(doc, element, left_indent_inches=0.2)
            elif element.name == "table":
                rows = element.find_all("tr")
                if not rows:
                    continue
                cols_count = max(len(row.find_all(["td", "th"])) for row in rows)
                table = doc.add_table(rows=len(rows), cols=cols_count)
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    cols = row.find_all(["td", "th"])
                    for c_idx, col in enumerate(cols):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = col.get_text().strip()

    @classmethod
    def convert_markdown_to_docx(
        cls,
        md_text: str,
        output_docx_path: Union[str, Path],
        ignore_frontmatter: bool = False,
    ) -> Path:
        """
        Converts markdown string into a styled Word document and writes to output_docx_path.
        """
        doc = docx.Document()
        cls.apply_default_styles(doc)

        fm_data, body_text = cls.extract_frontmatter(md_text)
        if not ignore_frontmatter and fm_data:
            cls.render_frontmatter(doc, fm_data)

        target_md = body_text if (ignore_frontmatter or fm_data) else md_text
        target_md = re.sub(r"~~(.*?)~~", r"<del>\1</del>", target_md)
        target_md = cls.preprocess_fenced_code_blocks(target_md)

        html = markdown.markdown(
            target_md, extensions=["fenced_code", "tables", "sane_lists"]
        )
        cls.html_to_docx(html, doc)

        out_path = Path(output_docx_path).resolve()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        return out_path

    @classmethod
    def convert_markdown_to_bytes(
        cls, md_text: str, ignore_frontmatter: bool = False
    ) -> bytes:
        """
        Converts markdown string into docx binary bytes in memory.
        """
        doc = docx.Document()
        cls.apply_default_styles(doc)

        fm_data, body_text = cls.extract_frontmatter(md_text)
        if not ignore_frontmatter and fm_data:
            cls.render_frontmatter(doc, fm_data)

        target_md = body_text if (ignore_frontmatter or fm_data) else md_text
        target_md = re.sub(r"~~(.*?)~~", r"<del>\1</del>", target_md)
        target_md = cls.preprocess_fenced_code_blocks(target_md)

        html = markdown.markdown(
            target_md, extensions=["fenced_code", "tables", "sane_lists"]
        )
        cls.html_to_docx(html, doc)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
