# /// script
# requires-python = ">=3.13"
# dependencies = ["python-docx", "markdown", "beautifulsoup4", "pyyaml"]
# ///

"""
version: 1.5.0
rev date: 8-11-2026
author: Will Morris
created date: 8-7-2026
usage: to_docx_processor.py <input_path|- > [output_path] [--nf|--no-frontmatter]
description: Converts markdown files or stdin input to docx format with professional styling,
  including frontmatter extraction, robust fenced/indented code block rendering inside lists & top-level,
  strikethrough text support, bumped heading sizes (+2pt), tables, and inline formatting.
"""

import re
import sys
from pathlib import Path
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import markdown  # type: ignore
from bs4 import BeautifulSoup, NavigableString
import yaml


def apply_default_styles(doc):
    """Applies a consistent professional style, font, and size to the document."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(51, 51, 51)  # Dark gray readability

    _format = style.paragraph_format
    _format.line_spacing = 1.15
    _format.space_after = Pt(6)


def set_codeblock_style(paragraph, fill_hex="F6F8FA", border_hex="D0D7DE"):
    """Adds light grey background shading and left accent border to a codeblock paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    pPr.append(shd)
    pbdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="{border_hex}"/></w:pBdr>'
    )
    pPr.append(pbdr)


def preprocess_fenced_code_blocks(md_text):
    """Preprocesses ```code``` blocks into explicit <pre><code> HTML blocks so Python-Markdown does not misinterpret them inside lists."""
    pattern = re.compile(
        r"^[ \t]*```(?P<lang>\w*)[ \t]*\r?\n(?P<code>.*?)\r?\n[ \t]*```[ \t]*$",
        re.MULTILINE | re.DOTALL,
    )

    def repl(match):
        lang = match.group("lang").strip()
        code = match.group("code")
        lang_class = f' class="language-{lang}"' if lang else ""
        return f"\n<pre><code{lang_class}>\n{code}\n</code></pre>\n"

    return pattern.sub(repl, md_text)


def extract_frontmatter(md_text):
    """Extracts YAML frontmatter from top of markdown document if present."""
    pattern = re.compile(r"^---\s*\r?\n(.*?)\r?\n---\s*\r?\n", re.DOTALL)
    match = pattern.match(md_text)
    if not match:
        return {}, md_text

    fm_raw = match.group(1)
    body_text = md_text[match.end():]
    try:
        fm_data = yaml.safe_load(fm_raw)
        if isinstance(fm_data, dict):
            return fm_data, body_text
    except Exception:
        pass
    return {}, md_text


def render_frontmatter(doc, fm_data):
    """Renders YAML frontmatter metadata as a styled header section at top of docx."""
    if not fm_data:
        return

    # Check for Title key
    title = (
        fm_data.pop("Title", None)
        or fm_data.pop("title", None)
        or fm_data.pop("Name", None)
        or fm_data.pop("name", None)
    )
    if title:
        p = doc.add_heading(str(title), level=1)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        for run in p.runs:
            run.font.size = Pt(20)

    # Render remaining metadata in a styled metadata summary table
    if fm_data:
        table = doc.add_table(rows=len(fm_data), cols=2)
        table.style = "Table Grid"
        table.autofit = True

        for idx, (k, v) in enumerate(fm_data.items()):
            row = table.rows[idx]
            cell_key = row.cells[0]
            cell_val = row.cells[1]

            cell_key.text = str(k).capitalize()
            cell_val.text = str(v)

            # Style key column bold
            p_k = cell_key.paragraphs[0]
            p_k.paragraph_format.space_after = Pt(2)
            if p_k.runs:
                p_k.runs[0].bold = True
                p_k.runs[0].font.size = Pt(9.5)

            p_v = cell_val.paragraphs[0]
            p_v.paragraph_format.space_after = Pt(2)
            if p_v.runs:
                p_v.runs[0].font.size = Pt(9.5)

            # Light shading for key background
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>')
            cell_key._tc.get_or_add_tcPr().append(shd)

        # Add space after metadata table
        doc.add_paragraph().paragraph_format.space_after = Pt(6)


def render_code_block(doc, element, left_indent_inches=0.2):
    """Renders a code block (<pre> or <code> block) as styled Consolas paragraphs with background shading."""
    code_text = element.get_text()

    # Clean leading language name if present as first line
    lines = code_text.splitlines()
    if lines and lines[0].strip() in [
        "python",
        "cmd",
        "bash",
        "sh",
        "json",
        "yaml",
        "yml",
        "ps1",
        "powershell",
        "html",
        "css",
        "javascript",
        "js",
        "ts",
    ]:
        lines = lines[1:]

    # Remove empty leading/trailing lines
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()

    if not lines:
        return

    for idx, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(2) if idx == 0 else Pt(0)
        p.paragraph_format.space_after = Pt(2) if idx == len(lines) - 1 else Pt(0)
        p.paragraph_format.left_indent = Inches(left_indent_inches)
        set_codeblock_style(p)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(36, 41, 46)


def add_inline_formatting(doc, paragraph, element):
    """Processes child nodes of an HTML element, preserving inline code, bold, italic, strikethrough, and links."""
    for child in element.contents:
        if isinstance(child, NavigableString):
            paragraph.add_run(str(child))
        elif child.name == "code":
            code_text = child.get_text()
            # If code element contains newlines or looks like a block, render as a code block paragraph
            if "\n" in code_text or code_text.startswith(
                ("python\n", "cmd\n", "bash\n", "# ", ":: ")
            ):
                render_code_block(doc, child, left_indent_inches=0.4)
            else:
                run = paragraph.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(
                    199, 37, 78
                )  # Pink/red for inline single-word code
        elif child.name in ["strong", "b"]:
            run = paragraph.add_run(child.get_text())
            run.bold = True
        elif child.name in ["em", "i"]:
            run = paragraph.add_run(child.get_text())
            run.italic = True
        elif child.name in ["del", "s", "strike"]:
            run = paragraph.add_run(child.get_text())
            run.font.strike = True
        elif child.name == "a":
            run = paragraph.add_run(child.get_text())
            run.font.color.rgb = RGBColor(3, 102, 214)  # Hyperlink blue
            run.underline = True
        elif child.name == "br":
            paragraph.add_run("\n")
        elif child.name == "pre":
            render_code_block(doc, child, left_indent_inches=0.4)
        elif child.name == "p":
            add_inline_formatting(doc, paragraph, child)
        else:
            paragraph.add_run(child.get_text())


def process_list_item(doc, li, list_style, indent_level=1):
    """Processes a single <li> element, correctly separating text, nested code blocks (<pre>/<code>), and sublists."""
    pres = li.find_all(["pre"], recursive=False)
    sublists = li.find_all(["ul", "ol"], recursive=False)

    # First add main list item paragraph for text before pre/sublists
    p = doc.add_paragraph(style=list_style)
    p.paragraph_format.left_indent = Inches(0.25 * indent_level)

    for child in li.contents:
        if isinstance(child, NavigableString):
            p.add_run(str(child))
        elif child.name == "p":
            add_inline_formatting(doc, p, child)
        elif child.name not in ["pre", "ul", "ol"]:
            soup_child = BeautifulSoup(str(child), "html.parser")
            target = soup_child.body if soup_child.body else soup_child
            add_inline_formatting(doc, p, target)

    # Render any code blocks nested inside the list item
    for pre in pres:
        render_code_block(doc, pre, left_indent_inches=0.25 * indent_level + 0.25)

    # Render any sublists nested inside the list item
    for sublist in sublists:
        sub_style = "List Bullet" if sublist.name == "ul" else "List Number"
        for sub_li in sublist.find_all("li", recursive=False):
            process_list_item(doc, sub_li, sub_style, indent_level + 1)


def html_to_docx(html_content, doc):
    """Parses HTML elements from markdown and renders them cleanly into the docx document."""
    soup = BeautifulSoup(html_content, "html.parser")

    for element in soup.children:
        if isinstance(element, NavigableString):
            continue

        if element.name == "h1":
            p = doc.add_heading(element.get_text(), level=1)
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            for run in p.runs:
                run.font.size = Pt(20)  # Bumped 2pt
        elif element.name == "h2":
            p = doc.add_heading(element.get_text(), level=2)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                run.font.size = Pt(16)  # Bumped 2pt
        elif element.name == "h3":
            p = doc.add_heading(element.get_text(), level=3)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.size = Pt(14)  # Bumped 2pt
        elif element.name == "h4":
            p = doc.add_heading(element.get_text(), level=4)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.size = Pt(12)  # Bumped 2pt
        elif element.name == "p":
            p = doc.add_paragraph()
            add_inline_formatting(doc, p, element)
        elif element.name in ["ul", "ol"]:
            list_style = "List Bullet" if element.name == "ul" else "List Number"
            for li in element.find_all("li", recursive=False):
                process_list_item(doc, li, list_style, indent_level=1)
        elif element.name == "blockquote":
            p = doc.add_paragraph(style="Intense Quote")
            add_inline_formatting(doc, p, element)
        elif element.name in ["pre", "div"] and (
            element.name == "pre" or "code" in element.get("class", [])
        ):
            render_code_block(doc, element, left_indent_inches=0.2)
        elif element.name == "table":
            rows = element.find_all("tr")
            if not rows:
                continue
            table = doc.add_table(
                rows=len(rows), cols=len(rows[0].find_all(["td", "th"]))
            )
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                cols = row.find_all(["td", "th"])
                for c_idx, col in enumerate(cols):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = col.get_text().strip()


def convert_text_to_docx(md_text, output_docx_path, ignore_frontmatter=False):
    """Converts raw markdown string to a styled docx file, handling frontmatter appropriately."""
    doc = docx.Document()
    apply_default_styles(doc)

    fm_data, body_text = extract_frontmatter(md_text)

    if not ignore_frontmatter and fm_data:
        render_frontmatter(doc, fm_data)

    target_md = body_text if (ignore_frontmatter or fm_data) else md_text

    # Preprocess markdown strikethrough ~~text~~ to <del>text</del>
    target_md = re.sub(r"~~(.*?)~~", r"<del>\1</del>", target_md)

    # Preprocess fenced code blocks into explicit <pre><code> HTML block elements
    target_md = preprocess_fenced_code_blocks(target_md)

    html = markdown.markdown(
        target_md, extensions=["fenced_code", "tables", "sane_lists"]
    )
    html_to_docx(html, doc)
    doc.save(output_docx_path)
    print(f"Docx generated successfully: {output_docx_path}")


def main():
    """Main CLI entrypoint handling input file, folder, stdin, and --nf / --no-frontmatter options."""
    args = [a for a in sys.argv[1:] if a not in ["--nf", "--no-frontmatter"]]
    ignore_frontmatter = "--nf" in sys.argv or "--no-frontmatter" in sys.argv

    if len(args) < 1:
        print(
            "Usage: to_docx_processor.py <input_path|- > [output_path] [--nf|--no-frontmatter]"
        )
        sys.exit(1)

    arg1 = args[0]
    custom_output = args[1] if len(args) > 1 and args[1] != "" else None

    # Handle stdin input
    if arg1 == "-" or not sys.stdin.isatty():
        md_text = sys.stdin.read()
        output_path = (
            Path(custom_output).resolve()
            if custom_output
            else Path("output.docx").resolve()
        )
        convert_text_to_docx(
            md_text, output_path, ignore_frontmatter=ignore_frontmatter
        )
        return

    input_path = Path(arg1).resolve()
    if not input_path.exists():
        print(f"Error: Path does not exist: {input_path}")
        sys.exit(1)

    if input_path.is_file():
        if input_path.suffix.lower() != ".md":
            print("Error: Provided file is not a markdown (.md) file.")
            sys.exit(1)

        output_path = (
            Path(custom_output).resolve()
            if custom_output
            else input_path.with_suffix(".docx")
        )
        with open(input_path, "r", encoding="utf-8") as f:
            md_text = f.read()
        convert_text_to_docx(
            md_text, output_path, ignore_frontmatter=ignore_frontmatter
        )

    elif input_path.is_dir():
        md_files = list(input_path.glob("*.md"))
        if not md_files:
            print(f"No .md files found in directory: {input_path}")
            sys.exit(0)

        for md_file in md_files:
            output_path = md_file.with_suffix(".docx")
            with open(md_file, "r", encoding="utf-8") as f:
                md_text = f.read()
            convert_text_to_docx(
                md_text, output_path, ignore_frontmatter=ignore_frontmatter
            )


if __name__ == "__main__":
    main()
